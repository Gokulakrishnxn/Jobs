"""
Build employment counts per NCO-2015 Family occupation from Census 2011 data.

Two real government sources, both India-total rows only (state=00, district=000):

1. census_b25a.xlsx - B-25A "Occupational classification of main workers
   other than cultivators and agricultural labourers by sex" - gives NCO-2004
   Family-level worker counts. Summed up to NCO-2015 Family codes via the
   crosswalk from build_crosswalk.py. Covers ~180.5M workers (household
   industry + "other" main workers).

2. census_pca.xlsx - Primary Census Abstract - gives India's fourfold worker
   classification (Cultivators / Agricultural Labourers / Household Industry
   / Other Workers). B-25A's universe is Household Industry + Other Workers,
   so Cultivators (95.9M) and Agricultural Labourers (86.2M) are added back
   as two aggregate categories that B-25A deliberately excludes - otherwise
   the treemap would be missing ~50% of India's workforce.

A handful of NCO-2004 codes (~14% of B-25A's workers) have no confident
match in the crosswalk (renamed/obsolete codes, or gaps in the PDF-derived
crosswalk); rather than silently dropping that employment, it's bucketed
into a synthetic "Unclassified" occupation per Division so total workforce
still reconciles to the real Census total.

Writes india/employment.json:
    {"by_family": {nco2015_family_code: total_persons}, ...}

Usage:
    uv run python india/build_employment.py
"""

import json
import openpyxl

DIVISION_NAMES = {
    "1": "Managers",
    "2": "Professionals",
    "3": "Technicians and Associate Professionals",
    "4": "Clerical Support Workers",
    "5": "Service and Sales Workers",
    "6": "Skilled Agricultural, Forestry and Fishery Workers",
    "7": "Craft and Related Trades Workers",
    "8": "Plant and Machine Operators and Assemblers",
    "9": "Elementary Occupations",
}


def load_india_rows(path):
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb["Sheet1"]
    rows = list(ws.iter_rows(values_only=True))
    # Data rows start after a fixed-height header block; find it by locating
    # the numbered header row ('1','2','3',...) rather than hardcoding index.
    start = None
    for i, r in enumerate(rows):
        if r and r[0] == "1" and r[1] == "2":
            start = i + 2  # skip the numbering row + one blank row
            break
    if start is None:
        raise RuntimeError(f"Could not find data start in {path}")
    return rows[start:]


def main():
    crosswalk = json.load(open("india/nco2004_to_2015.json"))

    rows = load_india_rows("india/raw/census_b25a.xlsx")
    by_family = {}
    unclassified_by_division = {}
    unmapped_examples = []

    for r in rows:
        if r is None or r[1] != "00" or r[2] != "000":
            continue  # keep only the India-total rows (not states/districts)
        division, family_2004, persons = r[4], r[7], r[9]
        if family_2004 == "0000" or persons is None:
            continue  # roll-up row, not a leaf occupation
        if not str(family_2004).isdigit():
            # "X999" (not reporting) / "X921" (unclassifiable) sentinel codes
            unclassified_by_division[division] = unclassified_by_division.get(division, 0) + persons
            continue

        targets = crosswalk.get(family_2004)
        if not targets:
            unclassified_by_division[division] = unclassified_by_division.get(division, 0) + persons
            if len(unmapped_examples) < 30:
                unmapped_examples.append((family_2004, r[8], persons))
            continue

        # Some NCO-2004 codes split across several NCO-2015 families (e.g.
        # doctors: NCO-2004 "2221" alone became NCO-2015's separate
        # Generalist/Specialist Medical Practitioners) - distribute this
        # row's persons across all of them by the crosswalk's weights
        # instead of assigning the whole count to a single winner.
        for family_2015, weight in targets.items():
            by_family[family_2015] = by_family.get(family_2015, 0) + persons * weight

    by_family = {code: round(v) for code, v in by_family.items()}
    b25a_total = sum(by_family.values()) + sum(unclassified_by_division.values())
    print(f"B-25A: {len(by_family)} NCO-2015 families populated, "
          f"{sum(by_family.values()):,} workers matched, "
          f"{sum(unclassified_by_division.values()):,} unclassified, "
          f"{b25a_total:,} total")

    # Add the two big agriculture aggregates B-25A excludes by design.
    pca_rows = load_india_rows_pca("india/raw/census_pca.xlsx")
    header, india_total = pca_rows
    pca = dict(zip(header, india_total))
    cultivators = pca["MAIN_CL_P"]
    ag_labourers = pca["MAIN_AL_P"]
    print(f"PCA: Cultivators {cultivators:,}, Agricultural Labourers {ag_labourers:,}")

    aggregates = {
        "CULT": {
            "code": "CULT",
            "title": "Cultivators",
            "division": DIVISION_NAMES["6"],
            "persons": cultivators,
            "description": (
                "Cultivators are persons who cultivate land they own or hold "
                "tenancy/lease rights to, or hold from the Government, "
                "bearing the risk of cultivation themselves. This is one of "
                "the Indian Census's four basic worker categories (Census of "
                "India 2011 Primary Census Abstract), not a National "
                "Classification of Occupations code — the Census does not "
                "further break cultivators down by NCO occupation."
            ),
        },
        "AGLAB": {
            "code": "AGLAB",
            "title": "Agricultural Labourers",
            "division": DIVISION_NAMES["9"],
            "persons": ag_labourers,
            "description": (
                "Agricultural labourers work on another person's land for "
                "wages in cash, kind, or a share of produce, without risk in "
                "cultivation and without owning/leasing the land they work. "
                "This is one of the Indian Census's four basic worker "
                "categories (Census of India 2011 Primary Census Abstract), "
                "not a National Classification of Occupations code — the "
                "Census does not further break agricultural labourers down "
                "by NCO occupation."
            ),
        },
    }

    unclassified = {}
    for division, persons in unclassified_by_division.items():
        code = f"UNCL{division}"
        if division == "X":
            title = "Occupation Not Reported / Unclassifiable"
            division_name = "Not Reported / Unclassifiable"
        else:
            title = f"Other {DIVISION_NAMES.get(division, division)} (unclassified)"
            division_name = DIVISION_NAMES.get(division, division)
        unclassified[code] = {
            "code": code,
            "title": title,
            "division": division_name,
            "persons": persons,
            "description": (
                "Workers whose Census 2011 usual-occupation response was not "
                "reported, or could not be classified into any NCO code, "
                "grouped here so total employment still reconciles to the "
                "real Census total."
            ) if division == "X" else (
                "A residual bucket of workers whose Census 2011 occupation "
                "code (NCO-2004) could not be confidently matched to a "
                "specific NCO-2015 occupation for this visualization — "
                "either a renamed/discontinued code, or a gap in the "
                "PDF-derived NCO-2004→NCO-2015 crosswalk. Grouped here "
                "so total employment still reconciles to the real Census "
                "total, at the cost of occupation-level detail for this "
                "slice."
            ),
        }

    output = {
        "by_family": by_family,
        "aggregates": aggregates,
        "unclassified": unclassified,
        "source": "Census of India 2011: B-25A (occupational, NCO-2004) + "
                   "Primary Census Abstract (cultivators/agricultural labourers)",
    }
    with open("india/employment.json", "w") as f:
        json.dump(output, f, indent=2)

    total = (sum(by_family.values()) + sum(v["persons"] for v in unclassified.values())
             + cultivators + ag_labourers)
    print(f"\nWrote india/employment.json")
    print(f"Total workers represented: {total:,}")
    print(f"(Census 2011 MAINWORK_P for comparison: {pca['MAINWORK_P']:,})")

    print("\nSample unmapped NCO-2004 codes (bucketed as unclassified):")
    for code, name, persons in unmapped_examples[:10]:
        print(f"  {code} {name}: {persons:,}")


def load_india_rows_pca(path):
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb["Sheet1"]
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    india_total = rows[1]  # first data row is India / Total
    assert india_total[0] == "00" and india_total[8] == "Total"
    return header, india_total


if __name__ == "__main__":
    main()
